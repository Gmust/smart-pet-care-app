from __future__ import annotations

import hashlib
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_PATH = (
    ROOT / "diploma" / "output" / "Smart_Pet_Care_CDV_Thesis_Working_Draft.docx"
)

EXPECTED_HEADINGS = [
    "Introduction",
    "1. Current state of knowledge",
    "2. Objective and scope",
    "3. Methodology",
    "4. Project implementation",
    "5. Deployment and operation",
    "Summary and conclusions",
    "Bibliography",
    "List of figures",
    "List of tables",
    "Streszczenie",
    "Abstract",
    "Team work-allocation sheet",
]
EXPECTED_BIBLIOGRAPHY_NUMBERS = list(range(1, 24))
TEMPORARY_CITATION_PATTERN = re.compile(
    r"\[(?:LIT|STD|SEC|EXPO|TSQ|MARKET)-\d+\]"
)
NUMERIC_CITATION_PATTERN = re.compile(r"\[(\d+(?:,\s*\d+)*)\]")
BIBLIOGRAPHY_ENTRY_PATTERN = re.compile(r"^\[(\d+)\]\s")


def close(actual: float, expected: float, tolerance: float = 0.02) -> bool:
    return abs(actual - expected) <= tolerance


def audit(path: Path) -> list[str]:
    errors: list[str] = []
    document = Document(path)

    if len(document.sections) != 2:
        errors.append(f"Expected 2 sections, found {len(document.sections)}")

    for index, section in enumerate(document.sections, start=1):
        if not close(section.page_width.cm, 21):
            errors.append(f"Section {index} width is not A4")
        if not close(section.page_height.cm, 29.7):
            errors.append(f"Section {index} height is not A4")
        for margin_name, margin in (
            ("top", section.top_margin),
            ("bottom", section.bottom_margin),
            ("left", section.left_margin),
            ("right", section.right_margin),
        ):
            if not close(margin.cm, 2.5):
                errors.append(
                    f"Section {index} {margin_name} margin is {margin.cm:.2f} cm"
                )

    normal = document.styles["Normal"]
    if normal.font.name != "Verdana":
        errors.append(f"Normal font is {normal.font.name!r}, not Verdana")
    if normal.font.size is None or not close(normal.font.size.pt, 10):
        errors.append("Normal font size is not 10 pt")
    if not close(float(normal.paragraph_format.line_spacing), 1.15):
        errors.append("Normal line spacing is not 1.15")
    first_line_indent = normal.paragraph_format.first_line_indent
    if first_line_indent is None or not close(first_line_indent.cm, 1.25):
        errors.append("Normal first-line indent is not 1.25 cm")

    heading_one = document.styles["Heading 1"]
    if heading_one.font.name != "Verdana":
        errors.append("Heading 1 font is not Verdana")
    if heading_one.font.size is None or not close(heading_one.font.size.pt, 12):
        errors.append("Heading 1 size is not 12 pt")
    if not heading_one.font.bold:
        errors.append("Heading 1 is not bold")
    if not heading_one.paragraph_format.page_break_before:
        errors.append("Heading 1 does not start on a new page")

    headings = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.style.name == "Heading 1"
    ]
    if headings != EXPECTED_HEADINGS:
        errors.append("Top-level headings are missing, reordered, or unexpected")

    all_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            paragraph.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        ]
    )
    for required_text in (
        "COLLEGIUM DA VINCI",
        "Faculty of Applied Sciences",
        "Major: INFORMATION TECHNOLOGY",
        "Engineering thesis",
        "Poznań [YEAR]",
        "frontend lead",
        "79 of 79 tests passed",
        "ignored generated declaration",
        "passing but not teardown-clean",
    ):
        if required_text not in all_text:
            errors.append(f"Required text is missing: {required_text}")

    placeholder_count = all_text.count("AUTHOR INPUT REQUIRED")
    if placeholder_count < 20:
        errors.append(
            f"Expected at least 20 author-input markers, found {placeholder_count}"
        )

    if len(document.tables) != 11:
        errors.append(f"Expected 11 tables, found {len(document.tables)}")

    temporary_citations = sorted(set(TEMPORARY_CITATION_PATTERN.findall(all_text)))
    if temporary_citations:
        errors.append(
            "Temporary research-register citations remain: "
            + ", ".join(temporary_citations)
        )

    paragraph_texts = [paragraph.text.strip() for paragraph in document.paragraphs]
    bibliography_index = paragraph_texts.index("Bibliography")
    list_of_figures_index = paragraph_texts.index("List of figures")
    bibliography_numbers = [
        int(match.group(1))
        for text in paragraph_texts[bibliography_index + 1 : list_of_figures_index]
        if (match := BIBLIOGRAPHY_ENTRY_PATTERN.match(text)) is not None
    ]
    if bibliography_numbers != EXPECTED_BIBLIOGRAPHY_NUMBERS:
        errors.append(
            "Bibliography entries are not the continuous ordered sequence 1-23"
        )

    non_bibliography_text = "\n".join(
        paragraph_texts[:bibliography_index]
        + paragraph_texts[list_of_figures_index:]
        + [
            paragraph.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        ]
    )
    cited_numbers = {
        int(number)
        for match in NUMERIC_CITATION_PATTERN.finditer(non_bibliography_text)
        for number in match.group(1).split(",")
    }
    expected_citations = set(EXPECTED_BIBLIOGRAPHY_NUMBERS)
    if cited_numbers != expected_citations:
        missing = sorted(expected_citations - cited_numbers)
        unexpected = sorted(cited_numbers - expected_citations)
        if missing:
            errors.append(
                "Bibliography entries without an in-text/table citation: "
                + ", ".join(str(number) for number in missing)
            )
        if unexpected:
            errors.append(
                "Numeric citations without a bibliography entry: "
                + ", ".join(str(number) for number in unexpected)
            )

    if len(document.inline_shapes) != 8:
        errors.append(
            f"Expected 8 inline thesis figures, found {len(document.inline_shapes)}"
        )

    captions = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Caption"
    ]
    if len(captions) != 19:
        errors.append(
            f"Expected 8 figure and 11 table captions, found {len(captions)}"
        )
    for caption_index, caption in enumerate(captions, start=1):
        if caption.alignment != WD_ALIGN_PARAGRAPH.LEFT:
            errors.append(f"Caption {caption_index} is not left-aligned")

    body_children = list(document.element.body)
    for table_index, child_index in enumerate(
        (
            index
            for index, child in enumerate(body_children)
            if child.tag == qn("w:tbl")
        ),
        start=1,
    ):
        if child_index == 0:
            errors.append(f"Table {table_index} has no preceding caption")
            continue
        preceding_xml = body_children[child_index - 1].xml
        if "SEQ Table" not in preceding_xml:
            errors.append(
                f"Table {table_index} is not immediately preceded by its caption"
            )

    non_caption_paragraph_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name != "Caption"
    )
    document_text = (
        "\n".join(paragraph.text for paragraph in document.paragraphs)
        + "\n"
        + "\n".join(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
    )
    for use_case_number in range(1, 19):
        use_case_id = f"UC-{use_case_number:02d}"
        if use_case_id not in document_text:
            errors.append(f"{use_case_id} is missing from the thesis traceability content")

    for table_reference in (
        "Table 1.1",
        "Table 2.1",
        "Table 2.2",
        "Table 2.3",
        "Table 2.4",
        "Table 2.5",
        "Table 3.1",
        "Table 4.1",
        "Table 4.2",
        "Table 5.1",
        "Table A.1",
    ):
        if table_reference not in non_caption_paragraph_text:
            errors.append(f"{table_reference} is not referenced from body prose")

    for figure_reference in (
        "Figure 3.1",
        "Figure 3.2",
        "Figure 3.3",
        "Figure 3.4",
        "Figure 3.5",
    ):
        if figure_reference not in non_caption_paragraph_text:
            errors.append(f"{figure_reference} is not referenced from body prose")
    if "Figures 3.6-3.8" not in non_caption_paragraph_text:
        errors.append("Figures 3.6-3.8 are not referenced from body prose")

    for figure_index, inline_shape in enumerate(document.inline_shapes, start=1):
        description = inline_shape._inline.docPr.get("descr")
        if description is None or not description.strip():
            errors.append(f"Figure {figure_index} has no alternative text")

    for table_index, table in enumerate(document.tables, start=1):
        if table.alignment != WD_TABLE_ALIGNMENT.CENTER:
            errors.append(f"Table {table_index} is not centre-aligned")
        for row_index, row in enumerate(table.rows, start=1):
            row_properties = row._tr.trPr
            if (
                row_properties is not None
                and row_properties.find(qn("w:trHeight")) is not None
            ):
                errors.append(
                    f"Table {table_index} row {row_index} has a fixed/declared height"
                )

    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
        footer_xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )

    if "TOC \\" not in document_xml:
        errors.append("No automatic TOC field found")
    if document_xml.count("SEQ Figure") != 8:
        errors.append("Expected 8 automatic Figure sequence fields")
    if document_xml.count("SEQ Table") != 11:
        errors.append("Expected 11 automatic Table sequence fields")
    if "updateFields" not in settings_xml:
        errors.append("Automatic field update setting is missing")
    if " PAGE " not in footer_xml:
        errors.append("No page-number field found in the body footer")

    return errors


def main() -> None:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_PATH
    errors = audit(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()

    if errors:
        print("FAIL")
        for error in errors:
            print(f"- {error}")
        raise SystemExit(1)

    document = Document(path)
    print("PASS")
    print(f"file: {path}")
    print(f"sha256: {digest}")
    print(f"sections: {len(document.sections)}")
    print(f"paragraphs: {len(document.paragraphs)}")
    print(f"tables: {len(document.tables)}")
    print(f"figures: {len(document.inline_shapes)}")
    print(
        "author_input_markers:",
        sum(
            paragraph.text.count("AUTHOR INPUT REQUIRED")
            for paragraph in document.paragraphs
        ),
    )


if __name__ == "__main__":
    main()
