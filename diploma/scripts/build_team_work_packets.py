from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "diploma" / "output" / "team-work-packets"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "222222"
MUTED = "666666"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(
    run,
    *,
    size: float = 11,
    bold: bool = False,
    italic: bool = False,
    color: str = INK,
) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def configure_document(document: Document, running_title: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run(f"SMART PET CARE  |  {running_title.upper()}")
    set_run_font(run, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.paragraph_format.space_before = Pt(0)
    run = footer_p.add_run("Working team packet  |  Page ")
    set_run_font(run, size=9, color=MUTED)
    add_page_field(footer_p)

    core = document.core_properties
    core.author = "Smart Pet Care diploma team"
    core.subject = "Team scope of work and evidence handoff"
    core.keywords = "Smart Pet Care, diploma, responsibilities, evidence, scope"


def add_title_block(
    document: Document,
    title: str,
    subtitle: str,
    owner: str,
    date: str = "26 July 2026",
) -> None:
    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    kicker.paragraph_format.space_after = Pt(6)
    run = kicker.add_run("SMART PET CARE DIPLOMA PROJECT")
    set_run_font(run, size=10, bold=True, color=BLUE)

    title_p = document.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run = title_p.add_run(title)
    set_run_font(run, size=23, bold=True, color=INK)

    subtitle_p = document.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(14)
    run = subtitle_p.add_run(subtitle)
    set_run_font(run, size=12.5, color=MUTED)

    values = [("Owner", owner), ("Status date", date)]
    for label, value in values:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}: ")
        set_run_font(label_run, size=10, bold=True, color=DARK_BLUE)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=10, color=INK)


def add_lead(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, size=11, bold=True, color=DARK_BLUE)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=11, color=INK)


def add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_run_font(run)


def add_numbered(document: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        run = paragraph.add_run(item)
        set_run_font(run)


def add_deliverables_table(document: Document, rows: list[tuple[str, str, str]]) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ("Deliverable", "Required content", "Completion evidence")
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, size=10, bold=True, color=WHITE)
        set_cell_shading(cell, BLUE)
    set_repeat_table_header(table.rows[0])

    for deliverable, content, evidence in rows:
        cells = table.add_row().cells
        for index, value in enumerate((deliverable, content, evidence)):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(value)
            set_run_font(run, size=9.5, bold=index == 0)
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[index], LIGHT_GRAY)
    set_table_geometry(table, [2100, 4140, 3120])


def add_checklist(document: Document, items: list[str]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = ""
    table.cell(0, 1).text = ""
    for index, text in enumerate(("Done", "Handoff item")):
        run = table.cell(0, index).paragraphs[0].add_run(text)
        set_run_font(run, size=10, bold=True, color=WHITE)
        set_cell_shading(table.cell(0, index), DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for item in items:
        cells = table.add_row().cells
        cells[0].text = ""
        cells[1].text = ""
        mark = cells[0].paragraphs[0].add_run("☐")
        set_run_font(mark, size=12, color=DARK_BLUE)
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cells[1].paragraphs[0].add_run(item)
        set_run_font(run, size=10)
    set_table_geometry(table, [780, 8580])


def add_document_boundary(document: Document) -> None:
    document.add_heading("Boundary of responsibility", level=2)
    add_lead(
        document,
        "Important",
        "Repository authorship supports responsibility but does not prove final "
        "acceptance. Do not mark work complete until the immutable revision, "
        "test result, and review evidence are attached.",
    )


def save_document(document: Document, filename: str) -> Path:
    path = OUTPUT_DIR / filename
    document.core_properties.title = path.stem.replace("_", " ")
    document.save(path)
    return path


def build_team_plan() -> Path:
    document = Document()
    configure_document(document, "Team work plan")
    add_title_block(
        document,
        "Team Work Plan and Handoff Rules",
        "Shared scope, deliverables, dependencies, and definition of done",
        "All four authors; Illia coordinates",
    )
    add_lead(
        document,
        "Objective",
        "Complete an evidence-based diploma thesis and a demonstrable first "
        "version of Smart Pet Care without presenting branch-only, untested, or "
        "version-2 work as a final result.",
    )

    document.add_heading("1. Confirmed responsibility map", level=1)
    add_deliverables_table(
        document,
        [
            (
                "Volodymyr Biletskyi",
                "Lead C# backend architecture, persistence, authentication, core "
                "APIs, reminders, notifications, infrastructure, and release evidence.",
                "Backend revision/tag, tests, deployment and migration records, "
                "reviewed thesis claims.",
            ),
            (
                "Kseniia Kushlak",
                "Backend features and research: photographs, feeding, weight "
                "history, C# classifier integration, literature and source evidence.",
                "Feature/test references, research log, source decisions, reviewed "
                "integration description.",
            ),
            (
                "Anastasia Leonova",
                "Product and idea leadership, frontend implementation, UX, care "
                "screens, screenshots, and usability/accessibility evidence.",
                "Accepted frontend revision, Figma/design history, screenshots, "
                "scenario and review results.",
            ),
            (
                "Illia Dolbnia",
                "Team and frontend leadership, AI service, system integration, "
                "quality gates, thesis assembly, and final handoff coordination.",
                "Mobile/AI revisions, model and deployment evidence, integrated "
                "test results, complete thesis package.",
            ),
        ],
    )

    document.add_heading("2. Shared scope decision", level=1)
    add_bullets(
        document,
        [
            "Use the author-confirmed first-version checklist as the drafting baseline.",
            "Treat checklist items 25-26 as version 2: the backend feeding capability "
            "may be described, but the completed mobile feeding experience is not a "
            "first-version result.",
            "Do not claim omitted items 27, 39-40, 42-43, 58-63, or 68 until the team "
            "classifies them.",
            "Use non-diagnostic language for health and assistant functionality.",
            "Name branch-only implementation as branch evidence until it is merged, "
            "tagged, or otherwise accepted.",
        ],
    )

    document.add_heading("3. Team workflow", level=1)
    add_numbered(
        document,
        [
            "Each owner reviews their individual packet and corrects factual errors.",
            "Each owner identifies one immutable revision for every repository or "
            "design artifact they claim.",
            "Each owner executes or supplies the required tests and preserves "
            "sanitized logs, screenshots, or reports.",
            "Each owner writes a short first-person explanation of important "
            "implementation decisions and problems encountered.",
            "All evidence is handed to Illia using stable paths or links; secrets and "
            "personal test data are removed.",
            "Illia integrates the evidence into the thesis and returns the affected "
            "sections to the owners for review.",
            "All four authors approve the responsibility table, scope, results, "
            "limitations, and official CDV contribution sheet.",
        ],
    )

    document.add_heading("4. Definition of done", level=1)
    add_checklist(
        document,
        [
            "The deliverable is linked to an immutable commit, tag, artifact digest, "
            "Figma version, or dated research record.",
            "The relevant tests or review scenarios have observed results.",
            "Failures and limitations are recorded, not hidden.",
            "The responsible author reviewed the corresponding thesis wording.",
            "No secret, token, private endpoint credential, or personal test data is "
            "included.",
            "The accepted mobile, backend, AI, and OpenAPI versions agree.",
            "Version-2 and unconfirmed items are clearly separated from first-version results.",
        ],
    )

    document.add_heading("5. Immediate team meeting decisions", level=1)
    add_checklist(
        document,
        [
            "Classify omitted checklist items 27, 39-40, 42-43, 58-63, and 68.",
            "Select the accepted mobile revision and decide the care-screen branch treatment.",
            "Select the accepted C# backend revision and deployment identifier.",
            "Confirm whether AI revision 74019ea was merged, tagged, or deployed.",
            "Select the final OpenAPI snapshot and regenerate the mobile client if required.",
            "Agree on the evidence handoff date and supervisor review date.",
        ],
    )
    return save_document(document, "01_Team_Work_Plan_and_Handoff.docx")


def build_volodymyr() -> Path:
    document = Document()
    configure_document(document, "Volodymyr scope")
    add_title_block(
        document,
        "Scope of Work - Volodymyr Biletskyi",
        "Lead C# backend developer and infrastructure owner",
        "Volodymyr Biletskyi (album 29934)",
    )
    add_lead(
        document,
        "Primary outcome",
        "Provide an accepted, testable, and documented C# backend release and "
        "verify every backend architecture, persistence, notification, and "
        "deployment claim attributed to you.",
    )

    document.add_heading("1. Owned technical scope", level=1)
    add_bullets(
        document,
        [
            "ASP.NET Core architecture, module boundaries, dependency registration, "
            "configuration, and OpenAPI exposure.",
            "PostgreSQL and Entity Framework Core model, migrations, relationships, "
            "ownership constraints, and data lifecycle.",
            "Authentication, authorization, users, pets, core CRUD, reminders, "
            "Firebase notification delivery, health, journal, and infrastructure.",
            "Backend release, containerization, environment configuration, diagnostics, "
            "and deployment or hosting evidence.",
            "Technical review of Kseniia's feature and chat-integration contributions "
            "where commits record co-authorship or shared ownership.",
        ],
    )

    document.add_heading("2. Required deliverables", level=1)
    add_deliverables_table(
        document,
        [
            (
                "Backend release baseline",
                "Final commit/tag, branch, build environment, API version, and relation "
                "to mobile OpenAPI snapshot.",
                "Commit URL or exported log, tag/image digest, clean build record.",
            ),
            (
                "Architecture description",
                "Modules, request flow, authentication, persistence, external "
                "integrations, and responsibility boundaries.",
                "Reviewed component/deployment diagrams and short author explanation.",
            ),
            (
                "Database evidence",
                "Database technology, final migration state, major relationships, "
                "backup/privacy limitations.",
                "Migration list, sanitized schema evidence, applied migration record.",
            ),
            (
                "Backend verification",
                "Unit/integration/API tests for owned modules, including negative and "
                "ownership cases.",
                "Exact commands, revision, environment, pass/fail totals, defect links.",
            ),
            (
                "Deployment evidence",
                "Backend host/container, configuration boundary, health check, logs, "
                "monitoring, and rollback approach.",
                "Sanitized deployment record, image digest, health response, log sample.",
            ),
            (
                "Thesis review",
                "Review Chapters 3.3, 3.4, 4.8, deployment material, results, and your "
                "contribution description.",
                "Dated approval or corrections with linked evidence.",
            ),
        ],
    )

    document.add_heading("3. Questions you must answer", level=1)
    add_numbered(
        document,
        [
            "Which backend revision is the final thesis release?",
            "Where and how is that revision deployed?",
            "Which PostgreSQL migration is current in the accepted environment?",
            "Which backend tests were executed on the accepted revision, and what failed?",
            "Which modules were mainly yours, which were shared, and which belong to Kseniia?",
            "What were the three most important backend problems and design decisions?",
        ],
    )

    add_document_boundary(document)
    add_bullets(
        document,
        [
            "Do not claim ownership of the Python classifier or condition-analysis logic.",
            "Do not treat an OpenAPI endpoint or Docker file as proof of successful runtime behavior.",
            "Do not include secrets, Firebase credentials, database passwords, or production personal data.",
        ],
    )

    document.add_heading("4. Send to Illia", level=1)
    add_checklist(
        document,
        [
            "Final backend revision/tag and repository link.",
            "Clean build and backend test report.",
            "Final OpenAPI JSON or URL with digest.",
            "Migration state and sanitized database architecture evidence.",
            "Deployment/image/health-check evidence.",
            "Reviewed architecture and deployment diagrams.",
            "First-person implementation decisions, problems, and limitations.",
            "Approval or corrections for your thesis responsibility description.",
        ],
    )
    return save_document(document, "02_Volodymyr_Biletskyi_Scope_of_Work.docx")


def build_kseniia() -> Path:
    document = Document()
    configure_document(document, "Kseniia scope")
    add_title_block(
        document,
        "Scope of Work - Kseniia Kushlak",
        "Backend feature developer and researcher",
        "Kseniia Kushlak (album 30081)",
    )
    add_lead(
        document,
        "Primary outcome",
        "Document and verify your backend features and C# assistant integration, "
        "and deliver a reproducible research package supporting the thesis's "
        "current-state-of-knowledge chapter.",
    )

    document.add_heading("1. Owned technical and research scope", level=1)
    add_bullets(
        document,
        [
            "Cloudinary-backed pet photograph functionality.",
            "Feeding backend API and related validation or persistence.",
            "Weight-history functionality and tests.",
            "C# classifier/chat integration, including request forwarding, chat "
            "persistence, bounded message history, failures, retry metadata, and "
            "resilience behavior.",
            "Literature search, source evaluation, market comparison support, and "
            "verification that thesis claims match the cited evidence.",
        ],
    )

    document.add_heading("2. Required deliverables", level=1)
    add_deliverables_table(
        document,
        [
            (
                "Backend feature evidence",
                "Exact owned modules, commits, shared/co-authored work, validation, "
                "storage, and API behavior.",
                "Commit links, module paths, tests, and short first-person explanation.",
            ),
            (
                "Classifier integration note",
                "C# ownership of authentication, pet/session checks, persistence, "
                "last-eight history, forwarding, errors, retry, metrics, and circuit breaker.",
                "Reviewed contract diagram and evidence from code/tests.",
            ),
            (
                "Feature verification",
                "Photo, feeding, weight, and chat integration happy-path and negative tests.",
                "Exact commands, revision, totals, failures, and resolved defects.",
            ),
            (
                "Research protocol",
                "Databases/search engines, queries, dates, inclusion/exclusion criteria, "
                "quality assessment, and duplicate handling.",
                "Research log and source-decision sheet.",
            ),
            (
                "Literature package",
                "Approved sources for pet-care information behavior, AI safety, mobile "
                "engineering, testing, privacy, and accessibility.",
                "Bibliographic metadata, PDFs/links, relevant page notes, claim mapping.",
            ),
            (
                "Thesis review",
                "Review Chapter 1, methodology research sections, backend feature claims, "
                "Section 4.8, limitations, and your contribution description.",
                "Dated approval or corrections.",
            ),
        ],
    )

    document.add_heading("3. Responsibility boundary", level=1)
    add_lead(
        document,
        "C# integration versus AI analysis",
        "Your classifier-integration work is the C# layer that stores messages and "
        "context, calls the Python service, and handles failures. Illia's Python AI "
        "service owns routing, condition prediction, urgency analysis, and generated "
        "informational responses.",
    )
    add_bullets(
        document,
        [
            "You may explain how the C# layer constructs and protects the request.",
            "Do not claim that the C# layer performs the medical-condition analysis.",
            "Do not claim clinical validation; document software behavior and safety limits.",
            "The feeding backend may be documented, but checklist items 25-26 place the "
            "mobile feeding interface and nutrition goals in version 2.",
        ],
    )

    document.add_heading("4. Questions you must answer", level=1)
    add_numbered(
        document,
        [
            "Which modules and commits are yours, and which were shared with Volodymyr?",
            "Which feature tests were executed on the accepted backend revision?",
            "What happens when the AI service times out, rate-limits, or returns invalid data?",
            "Which research sources should remain in the final thesis, and why?",
            "What sources or claims should be removed because evidence is too weak?",
            "What were your most important implementation or research decisions?",
        ],
    )

    document.add_heading("5. Send to Illia", level=1)
    add_checklist(
        document,
        [
            "Feature-to-commit/module ownership list.",
            "Photo, feeding, weight, and classifier-integration test report.",
            "Reviewed chat/classifier ownership description and diagram.",
            "Research protocol and dated search log.",
            "Approved bibliography metadata and source-to-claim notes.",
            "List of rejected or limited sources and reasons.",
            "First-person implementation/research problems and decisions.",
            "Approval or corrections for your thesis responsibility description.",
        ],
    )
    return save_document(document, "03_Kseniia_Kushlak_Scope_of_Work.docx")


def build_anastasia() -> Path:
    document = Document()
    configure_document(document, "Anastasia scope")
    add_title_block(
        document,
        "Scope of Work - Anastasia Leonova",
        "Idea/product lead and frontend developer",
        "Anastasia Leonova (album 29945)",
    )
    add_lead(
        document,
        "Primary outcome",
        "Provide evidence for the product concept and deliver an accepted, "
        "backend-integrated frontend scope with screenshots and usability or "
        "accessibility evaluation suitable for the thesis.",
    )

    document.add_heading("1. Owned product and frontend scope", level=1)
    add_bullets(
        document,
        [
            "Product idea, target-user problem, user journeys, feature priorities, and "
            "first-version versus future-version decisions.",
            "Care-related frontend branch and assigned pages, components, forms, states, "
            "and interaction flows.",
            "Visual and interaction design evidence, including Figma versions or "
            "equivalent dated design artifacts.",
            "Representative screenshots and user-operation instructions.",
            "Usability, accessibility, localization, and product-consistency review.",
        ],
    )

    document.add_heading("2. Required deliverables", level=1)
    add_deliverables_table(
        document,
        [
            (
                "Product concept evidence",
                "Problem statement, users, needs, original idea decisions, roadmap, and "
                "why features were included, deferred, or rejected.",
                "Dated Confluence/Figma/meeting history and first-person explanation.",
            ),
            (
                "Accepted frontend revision",
                "Final care/frontend branch treatment, merged revision or explicit "
                "prototype status, backend integration, and removal/approval of mocks.",
                "Commit/PR reference, clean build, route list, integration result.",
            ),
            (
                "Design evidence",
                "Screen flows, components, states, accessibility annotations, and major "
                "design alternatives.",
                "Figma version links/exports and decision notes.",
            ),
            (
                "Screenshot package",
                "Sanitized first-version screens covering core user journeys and "
                "important empty/loading/error states.",
                "Named image files mapped to thesis sections and final build/revision.",
            ),
            (
                "UX/accessibility review",
                "Representative tasks, participant/reviewer context, device, focus, "
                "labels, scaling, contrast, touch targets, errors, and observations.",
                "Completed scenario sheet with findings, defects, and limitations.",
            ),
            (
                "Thesis review",
                "Review product scope, use cases, frontend implementation, screenshots, "
                "future work, and your contribution description.",
                "Dated approval or corrections.",
            ),
        ],
    )

    document.add_heading("3. Immediate frontend decisions", level=1)
    add_checklist(
        document,
        [
            "Confirm which care-screen branch revision belongs in the evaluated release.",
            "Identify every mock data source and replace it with the accepted backend "
            "integration or explicitly classify it as prototype-only.",
            "Confirm that checklist items 25-26 remain version 2.",
            "Help classify omitted items 27, 39-40, 42-43, 58-63, and 68.",
            "Confirm whether Polish localization is future work; the inspected baseline "
            "contains English resources only.",
            "Map final screenshots to the exact accepted Android build.",
        ],
    )

    document.add_heading("4. Questions you must answer", level=1)
    add_numbered(
        document,
        [
            "Which product decisions and user flows were originally yours?",
            "Which frontend files, screens, components, or branches did you implement?",
            "Which care features are merged and backend-integrated today?",
            "Which features remain prototype, mock-backed, unconfirmed, or version 2?",
            "What design alternatives were considered and why were they rejected?",
            "What usability or accessibility problems were found and fixed?",
        ],
    )

    add_document_boundary(document)
    add_bullets(
        document,
        [
            "A Figma frame or branch proves work exists, not that it is in the accepted release.",
            "Screenshots must not contain real user names, email addresses, tokens, or private pet-health data.",
            "Do not describe visual consistency as a usability result without an observed task or review.",
        ],
    )

    document.add_heading("5. Send to Illia", level=1)
    add_checklist(
        document,
        [
            "Product-concept authorship and decision evidence.",
            "Final frontend/care revision or explicit prototype classification.",
            "Owned screen/component/file list.",
            "Figma/design versions and key decision notes.",
            "Sanitized screenshot package mapped to the final build.",
            "Usability/accessibility scenario results and defect list.",
            "First-person product/frontend problems and decisions.",
            "Approval or corrections for your thesis responsibility description.",
        ],
    )
    return save_document(document, "04_Anastasia_Leonova_Scope_of_Work.docx")


def build_illia() -> Path:
    document = Document()
    configure_document(document, "Illia scope")
    add_title_block(
        document,
        "Scope of Work - Illia Dolbnia",
        "Team lead, frontend lead, and AI backend-service lead",
        "Illia Dolbnia (album 30151)",
    )
    add_lead(
        document,
        "Primary outcome",
        "Freeze and integrate the accepted mobile, backend, AI, and contract "
        "revisions; verify the AI-service behavior and mobile quality gates; and "
        "assemble the complete evidence-based thesis package.",
    )

    document.add_heading("1. Owned leadership and technical scope", level=1)
    add_bullets(
        document,
        [
            "Team scope, schedule, dependencies, review flow, evidence register, and "
            "university/supervisor coordination.",
            "Mobile-client architecture, shared frontend integration, authentication, "
            "pets, reminders, notifications, assistant client, recovery, and quality gates.",
            "Python AI microservice: unified `/chat`, condition classifier, deterministic "
            "triage safety, Gemini-generated text boundary, tests, CI, model lifecycle, "
            "and AWS deployment definition.",
            "End-to-end integration between the mobile client, C# backend, and AI service.",
            "Thesis assembly, consistency, diagrams, AI-use disclosure, final QA, and "
            "defence preparation.",
        ],
    )

    document.add_heading("2. Required deliverables", level=1)
    add_deliverables_table(
        document,
        [
            (
                "Scope freeze",
                "Classified checklist, accepted/rejected/version-2 items, final "
                "requirements, and accepted mobile/backend/AI/OpenAPI revisions.",
                "Signed team decision record and immutable revision list.",
            ),
            (
                "Mobile release evidence",
                "Clean mobile revision, generated client, lint, typecheck, tests, Android "
                "build/install, critical scenarios, and screenshots.",
                "Commands, logs, artifact digest, device details, results, defects.",
            ),
            (
                "AI release evidence",
                "Decision on feature revision 74019ea, model/backend version, `/chat` "
                "contract, safety cases, evaluation data, CI, deployment, monitoring.",
                "Merged/tagged/deployed revision, model card/report, test and live records.",
            ),
            (
                "Integration evidence",
                "Mobile-to-C# and C#-to-Python request/response flows, auth, history, "
                "timeouts, retry, failures, and recovery.",
                "Sanitized live exchanges and mapped end-to-end scenarios.",
            ),
            (
                "Thesis package",
                "All required chapters, contribution evidence, figures, tables, "
                "bibliography, appendices, screenshots, results, limitations, and AI disclosure.",
                "Reviewed DOCX/PDF, completed markers, signatures, supervisor feedback.",
            ),
            (
                "Team coordination",
                "Evidence intake, owner reviews, issue closure, deadlines, and defence split.",
                "Dated checklist, review approvals, and final submission record.",
            ),
        ],
    )

    document.add_heading("3. Immediate technical priorities", level=1)
    add_numbered(
        document,
        [
            "Decide whether AI feature revision 74019ea is merged, tagged, or deployed; "
            "do not leave the thesis pointing to a branch without status.",
            "Select the final C# backend and OpenAPI versions with Volodymyr and Kseniia.",
            "Select the final mobile revision with Anastasia and decide the care-screen "
            "branch treatment.",
            "Reconcile the `/chat` contract between the C# backend and Python service.",
            "Run the AI unit, regression, safety, contract, model, CI, and deployment checks.",
            "Run mobile lint, typecheck, tests, clean Android build/install, and mapped "
            "device scenarios.",
            "Collect every owner's evidence and incorporate only reviewed claims.",
        ],
    )

    document.add_heading("4. AI-service evaluation minimum", level=1)
    add_checklist(
        document,
        [
            "Exact code revision, model release, dataset/split identifiers, and configuration recorded.",
            "General-care, health, emergency, clarification, abstention, malformed-input, "
            "provider-failure, and low-confidence cases executed.",
            "Expected and observed output recorded without claiming clinical validation.",
            "Deterministic safety behavior tested independently of Gemini availability.",
            "No secret keys, raw personal messages, or private health data retained in evidence.",
            "Latency, failure, and deployment observations include environment and sample size.",
            "Known limitations and model/provider dependencies are explicit.",
        ],
    )

    document.add_heading("5. Questions you must close", level=1)
    add_numbered(
        document,
        [
            "What exact artifact is the evaluated first version?",
            "Which omitted checklist items are included, deferred, or rejected?",
            "Is AI revision 74019ea the deployed service? If not, what is?",
            "Which assistant/model results can be supported without veterinary validation?",
            "Which Android scenarios passed, failed, were blocked, or were not run?",
            "Has every author approved the claims and contribution description attributed to them?",
            "Has the supervisor approved the title, objective, scope, thesis language, and final package?",
        ],
    )

    document.add_heading("6. Final coordination checklist", level=1)
    add_checklist(
        document,
        [
            "Receive and review Volodymyr's complete backend packet.",
            "Receive and review Kseniia's backend/research packet.",
            "Receive and review Anastasia's product/frontend packet.",
            "Freeze accepted mobile, backend, AI, and contract revisions.",
            "Execute or verify all mandatory tests and scenarios.",
            "Replace every remaining thesis author-input marker.",
            "Update Word fields, inspect all pages, and export the final PDF.",
            "Collect all signatures and supervisor approval.",
            "Prepare role-specific defence sections and demonstration backup.",
        ],
    )
    return save_document(document, "05_Illia_Dolbnia_Scope_of_Work.docx")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    paths = [
        build_team_plan(),
        build_volodymyr(),
        build_kseniia(),
        build_anastasia(),
        build_illia(),
    ]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
